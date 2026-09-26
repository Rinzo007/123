"""Лист «Маршруты» XLSX-экспорта."""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from typing import Any

from ..models import RouteData
from ..support import direction_endpoints, fmt_curv_cell, source_label
from .helpers import cached_type_label, finish_sheet, safe_append

_ROUTES_CORE_HEADERS = [
    "Город", "Тип", "Номер", "ID", "Направлений",
]

_ROUTES_TAIL_HEADERS = [
    "№ напр.",
    "Направление", "От", "До", "Длина, км", "Коэфф. непрямолинейности",
    "Остановок", "Ссылка",
]

_ROUTES_NUM_FMT_BY_HEADER = {
    "Длина, км": "0.0",
    "Коэфф. непрямолинейности": "0.00",
    "Остановок": "0",
    "Площадь объектов, тыс. м²": "#,##0.0",
    "Население, чел.": "#,##0",
    "Плотность, чел/км²": "#,##0.0",
    "Точек POI": "0",
    "Суммарное value": "#,##0.0",
}

_ROUTES_SOURCE_WIDTH = 12
_ROUTES_ACTIVE_WIDTH = 10


def routes_headers(*, overture_stats=None, poi_stats=None,
                    poi_stops_stats=None, show_source: bool = True) -> list[str]:
    headers = list(_ROUTES_CORE_HEADERS)
    if show_source:
        headers.append("Источник")
    headers.append("Активен")
    headers += _ROUTES_TAIL_HEADERS
    if overture_stats:
        headers += ["Площадь объектов, тыс. м²"]
    if poi_stats:
        headers += ["Точек POI", "Суммарное value"]
    if poi_stops_stats:
        headers += ["POI у остановок"]
    return headers


def _stat_cell(st, *, attr: str, scale: float = 1.0, default=None):
    """Значение ячейки из объекта статистики направления (или ``default``)."""
    if st is None:
        return default
    return round(getattr(st, attr) / scale, 1)


def route_direction_row(rd: RouteData, di: int, direction: Any, base_row: list[Any], *,
                        overture_stats=None,
                        overture_dir_stats=None, poi_stats=None, poi_dir_stats=None,
                        poi_stops_stats=None, poi_stops_dir_stats=None) -> list[Any]:
    first, last = direction_endpoints(direction)
    row = [*base_row, di + 1, direction.name or f"{first} → {last}", first, last,
           direction.km, fmt_curv_cell(direction.curvilinearity), len(direction.stops), rd.url]
    row += _stat_row(rd, di, overture_dir_stats if overture_stats else None,
                     attr="total_area_m2", scale=1e3)
    if poi_stats:
        st = (poi_dir_stats or {}).get((rd.route_id, di))
        row += [st.count if st else None, round(st.total_value, 1) if st else None]
    row += _stat_row(rd, di, poi_stops_dir_stats if poi_stops_stats else None,
                     attr="count", scale=1.0)
    return row


def _stat_row(rd: RouteData, di: int, stats, *, attr: str, scale: float, fallback=None) -> list[Any]:
    """Возвращает ячейку со статистикой направления (пусто, если stats отключены)."""
    if stats is None:
        return []
    st = stats.get((rd.route_id, di)) or fallback
    return [_stat_cell(st, attr=attr, scale=scale)]


def _style_row(ws: Any, num_fmt: Mapping[int, str], wrap_cols: list[int], link_col: int | None) -> None:
    from openpyxl.styles import Alignment, Font
    row = ws.max_row
    for col, fmt in num_fmt.items():
        ws.cell(row=row, column=col).number_format = fmt
    for col in wrap_cols:
        ws.cell(row=row, column=col).alignment = Alignment(wrap_text=True, vertical="top")
    if link_col:
        cell = ws.cell(row=row, column=link_col)
        if cell.value:
            cell.hyperlink = str(cell.value)
            cell.font = Font(color="0563C1", underline="single")


def write_routes_sheet(wb: Any, *, city: str, kml_routes: Sequence[RouteData],
                       overture_stats=None,
                       overture_dir_stats=None, poi_stats=None, poi_dir_stats=None,
                       poi_stops_stats=None, poi_stops_dir_stats=None) -> None:
    sources = {route.source for route in kml_routes}
    show_source = len(sources) > 1
    headers = routes_headers(overture_stats=overture_stats, poi_stats=poi_stats,
                             poi_stops_stats=poi_stops_stats,
                             show_source=show_source)
    link_col = headers.index("Ссылка") + 1 if "Ссылка" in headers else None
    num_fmt = {i: fmt for i, h in enumerate(headers, 1) if (fmt := _ROUTES_NUM_FMT_BY_HEADER.get(h))}
    wrap_cols = [i for i, h in enumerate(headers, 1) if h in ("Направление", "От", "До")]
    ws = wb.create_sheet("Маршруты")
    safe_append(ws, headers)

    route_label_cache: dict[int, str] = {}
    route_source_cache: dict[int, str] = {}
    for route in kml_routes:
        route_label_cache[route.route_id] = cached_type_label(route.route_type)
        route_source_cache[route.route_id] = source_label(route.source)

    for route in sorted(kml_routes, key=lambda r: (getattr(r.route_type, "value", r.route_type), r.name)):
        base_row = [city, route_label_cache[route.route_id], route.name, route.route_id,
                    len(route.directions)]
        if show_source:
            base_row.append(route_source_cache[route.route_id])
        base_row.append("Да" if route.active else "Нет")
        for di, direction in enumerate(route.directions):
            safe_append(ws, route_direction_row(route, di, direction, base_row,
                                           overture_stats=overture_stats, overture_dir_stats=overture_dir_stats,
                                           poi_stats=poi_stats, poi_dir_stats=poi_dir_stats,
                                           poi_stops_stats=poi_stops_stats, poi_stops_dir_stats=poi_stops_dir_stats))
            if num_fmt or wrap_cols or link_col:
                _style_row(ws, num_fmt, wrap_cols, link_col)

    widths = [12, 10, 12, 8, 10]  # Город, Тип, Номер, ID, Направлений
    if show_source:
        widths.append(_ROUTES_SOURCE_WIDTH)  # Источник
    widths.append(_ROUTES_ACTIVE_WIDTH)  # Активен
    widths += [8, 20, 15, 15, 10, 10, 10, 30]  # № напр. → Ссылка
    if len(headers) > len(widths):
        widths += [12] * (len(headers) - len(widths))
    finish_sheet(ws, headers, num_fmt=num_fmt, wrap=set(wrap_cols), link_col=link_col,
                 fixed_widths=widths, inline_styled=True)
