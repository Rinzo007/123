"""Общие помощники XLSX-экспорта.

Модуль не содержит логики конкретных листов; здесь находятся только
переиспользуемые расчёты и оформление worksheet.
"""

from __future__ import annotations

import logging
import math
import unicodedata  # <-- added import
from collections.abc import Mapping, Sequence
from functools import lru_cache
from typing import Any

from ..support import type_label

logger = logging.getLogger("wikiroutes.xlsx.helpers")

# Символы, с которых Excel может начать интерпретировать значение как формулу.
_FORMULA_PREFIXES: tuple[str, ...] = ("=", "+", "-", "@", "\t", "\r")


def safe_append(ws: Any, row: Sequence[object]) -> None:
    """Записывает строку в лист, удаляя недопустимые символы и защищая от formula injection."""
    # --- Очистка каждого строкового значения от управляющих символов (кроме \t,\n,\r) ---
    sanitized_row = []
    for value in row:
        if isinstance(value, str):
            cleaned = ''.join(
                ch for ch in value
                if unicodedata.category(ch) != 'Cc' or ch in {'\t', '\n', '\r'}
            )
            sanitized_row.append(cleaned)
        else:
            sanitized_row.append(value)

    ws.append(sanitized_row)
    last_row = ws.max_row

    # --- Предотвращение formula injection для строк, начинающихся с опасных символов ---
    for column, value in enumerate(sanitized_row, start=1):
        if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
            ws.cell(row=last_row, column=column).data_type = "s"


@lru_cache(maxsize=128)
def cached_type_label(route_type: object) -> str:
    """Кэширует преобразование типа транспорта в русскую метку."""
    return type_label(route_type)


def network_density(
    net_metrics: Mapping[str, Any] | None,
    bbox: Any | None,  # может быть кортеж или объект BBox
) -> Any:
    """Плотность уникальной сети, км/км², по bbox города."""
    if not net_metrics or not bbox:
        return ""

    # Извлечение координат из BBox или кортежа
    if hasattr(bbox, "min_lat"):
        min_lat = bbox.min_lat
        min_lon = bbox.min_lon
        max_lat = bbox.max_lat
        max_lon = bbox.max_lon
    else:
        try:
            min_lat, min_lon, max_lat, max_lon = bbox
        except (TypeError, ValueError):
            return ""

    lat_c = (min_lat + max_lat) / 2
    area_km2 = ((max_lat - min_lat) * 111.0) * (
        (max_lon - min_lon) * 111.0 * math.cos(math.radians(lat_c))
    )

    if area_km2 <= 0:
        return ""

    return round(net_metrics.get("unique_km", 0.0) / area_km2, 3)


def _style_header_row(ws: Any) -> None:
    """Закрашивает шапку и центрирует заголовки."""
    from openpyxl.styles import Alignment, Font, PatternFill

    head_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    head_font = Font(color="FFFFFF", bold=True)
    for c in ws[1]:
        c.fill, c.font = head_fill, head_font
        c.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )


def _apply_fixed_widths(ws: Any, fixed_widths: Sequence[int]) -> None:
    """Задаёт фиксированные ширины колонок."""
    from openpyxl.utils import get_column_letter

    for i, width in enumerate(fixed_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width


def _apply_auto_widths(ws: Any, headers: Sequence[object]) -> None:
    """Подбирает ширины колонок по содержимому (не более 62 символов)."""
    from openpyxl.utils import get_column_letter

    max_rows = min(ws.max_row, 200)
    for i, header in enumerate(headers, 1):
        width = len(str(header))
        for row in ws.iter_rows(
            min_row=2,
            max_row=max_rows,
            min_col=i,
            max_col=i,
        ):
            value = row[0].value
            if value is not None:
                width = max(
                    width,
                    min(max(len(line) for line in str(value).split("\n")), 60),
                )
        ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 62)


def _apply_number_formats(ws: Any, num_fmt: Mapping[int, str]) -> None:
    """Применяет числовые форматы по колонкам."""
    for col, fmt in num_fmt.items():
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            row[0].number_format = fmt


def _apply_wrap_alignment(ws: Any, wrap: set[int]) -> None:
    """Включает перенос текста для указанных колонок."""
    from openpyxl.styles import Alignment

    for col in wrap:
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            row[0].alignment = Alignment(wrap_text=True, vertical="top")


def _apply_link_column(ws: Any, link_col: int, link_font: Any) -> None:
    """Превращает значения колонки в гиперссылки."""
    for row in ws.iter_rows(min_row=2, min_col=link_col, max_col=link_col):
        if row[0].value:
            row[0].hyperlink = str(row[0].value)
            row[0].font = link_font


def finish_sheet(
    ws: Any,
    headers: Sequence[object],
    num_fmt: Mapping[int, str] | None = None,
    wrap: set[int] | None = None,
    link_col: int | None = None,
    fixed_widths: Sequence[int] | None = None,
    inline_styled: bool = False,
) -> None:
    """Оформляет worksheet: шапка, ширины, форматы, ссылки и перенос."""
    from openpyxl.styles import Font

    link_font = Font(color="0563C1", underline="single")

    _style_header_row(ws)

    ws.freeze_panes = "A2"

    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions

    if fixed_widths:
        _apply_fixed_widths(ws, fixed_widths)
    else:
        _apply_auto_widths(ws, headers)

    if num_fmt and not inline_styled:
        _apply_number_formats(ws, num_fmt)

    if wrap and not inline_styled:
        _apply_wrap_alignment(ws, wrap)

    if link_col and not inline_styled:
        _apply_link_column(ws, link_col, link_font)


__all__ = ["cached_type_label", "finish_sheet", "network_density", "safe_append"]